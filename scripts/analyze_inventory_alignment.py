import json
import re
import sys
import zipfile
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
EXTRACT_DIR = ROOT / ".tmp_inventory_extract_20260831"
SOURCE_XLSX = EXTRACT_DIR / "inventory_source.xlsx"
STORE_JSON = ROOT / "server" / "data" / "store.json"
REPORT_MD = ROOT / "docs" / "inventory-alignment-report.md"


def text(value):
    return str(value).strip() if value is not None else ""


def number(value):
    if value is None or value == "":
        return 0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0


def normalize(value):
    value = text(value).lower()
    value = re.sub(r"\s+", "", value)
    value = value.replace("（", "(").replace("）", ")")
    value = value.replace("-", "").replace("_", "")
    return value


def classify(name):
    n = normalize(name)
    if any(word in n for word in ["维修", "换屏", "换电池", "主板", "进水", "检测"]):
        return "维修"
    if any(word in n for word in ["号卡", "卡利润", "元号卡", "套餐卡", "流量卡"]) or re.search(r"卡\d+元档", n):
        return "号卡"
    accessory_words = ["充电", "数据线", "套充", "耳机", "钢化膜", "手机壳", "保护套", "支架", "充电宝"]
    phone_brands = ["iphone", "iqoo", "荣耀", "红米", "小米", "vivo", "oppo", "realme", "一加", "nova"]
    phone_spec = re.search(r"\d+\+\d+", n) or re.search(r"(128|256|512)g", n)
    if not any(word in n for word in accessory_words) and (any(word in n for word in phone_brands) or phone_spec):
        return "手机"
    return "零配件"


def read_source_rows():
    workbook = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = None
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if text(row[0]) == "行号" and text(row[2]) == "商品编号":
            header_row = index
            break
    if not header_row:
        raise RuntimeError("没有找到库存表头")

    rows = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not text(row[2]) and not text(row[3]):
            continue
        qty = number(row[10])
        if qty <= 0:
            continue
        rows.append(
            {
                "code": text(row[2]),
                "name": text(row[3]),
                "shortName": text(row[4]),
                "spec": text(row[5]),
                "model": text(row[6]),
                "unit": text(row[8]) or "件",
                "costPrice": number(row[9]),
                "quantity": qty,
                "avgCost": number(row[12]),
                "stockValue": number(row[13]),
                "retailPrice": number(row[14]),
                "presetPrice": number(row[15]),
                "type": classify(row[3]),
            }
        )
    return rows


def read_current_rows():
    store = json.loads(STORE_JSON.read_text(encoding="utf-8"))
    products = {item["skuId"]: item for item in store["products"]}
    qty_by_sku = defaultdict(float)
    serials_by_sku = defaultdict(list)
    for stock in store["inventory"]:
        if stock.get("warehouseId") and stock.get("warehouseId") != "wh-sale":
            continue
        qty_by_sku[stock["skuId"]] += number(stock.get("quantity"))
        serials_by_sku[stock["skuId"]].extend(stock.get("serials") or [])

    rows = []
    for product in store["products"]:
        stock_managed = product.get("stockManaged", True) is not False
        rows.append(
            {
                "skuId": product["skuId"],
                "name": product["name"],
                "type": product.get("type", ""),
                "unit": product.get("unit", "件"),
                "costPrice": number(product.get("costPrice")),
                "salePrice": number(product.get("salePrice")),
                "quantity": qty_by_sku[product["skuId"]] if stock_managed else 0,
                "stockManaged": stock_managed,
                "serialCount": len(serials_by_sku[product["skuId"]]),
            }
        )
    return rows


def similarity(a, b):
    a_norm = normalize(a)
    b_norm = normalize(b)
    if not a_norm or not b_norm:
        return 0
    if a_norm == b_norm:
        return 1
    if a_norm in b_norm or b_norm in a_norm:
        return 0.92
    return SequenceMatcher(None, a_norm, b_norm).ratio()


def best_match(row, candidates):
    scored = [(similarity(row["name"], item["name"]), item) for item in candidates]
    scored.sort(key=lambda item: item[0], reverse=True)
    return scored[0] if scored else (0, None)


def money(value):
    if abs(value - round(value)) < 0.00001:
        return str(int(round(value)))
    return f"{value:.2f}"


def qty(value):
    if abs(value - round(value)) < 0.00001:
        return str(int(round(value)))
    return f"{value:.2f}"


def table(rows, headers):
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(row.get(header, "")).replace("|", "/") for header in headers) + " |")
    return "\n".join(lines)


def main():
    source_rows = read_source_rows()
    current_rows = read_current_rows()
    stock_source_rows = [row for row in source_rows if row["type"] != "维修"]
    source_by_type = defaultdict(int)
    source_qty_by_type = defaultdict(float)
    for row in source_rows:
        source_by_type[row["type"]] += 1
        source_qty_by_type[row["type"]] += row["quantity"]

    current_by_type = defaultdict(int)
    current_qty_by_type = defaultdict(float)
    for row in current_rows:
        current_by_type[row["type"]] += 1
        current_qty_by_type[row["type"]] += row["quantity"]

    matches = []
    weak_matches = []
    for current in current_rows:
        score, source = best_match(current, source_rows)
        if score >= 0.9:
            matches.append((current, source, score))
        else:
            weak_matches.append((current, source, score))

    matched_source_codes = {source["code"] for _, source, _ in matches}
    source_missing_in_current = [row for row in stock_source_rows if row["code"] not in matched_source_codes]

    conflicts = []
    for current, source, score in matches:
        if current["stockManaged"]:
            qty_diff = current["quantity"] - source["quantity"]
        else:
            qty_diff = 0
        cost_diff = current["costPrice"] - source["avgCost"]
        price_diff = current["salePrice"] - (source["presetPrice"] or source["retailPrice"])
        has_conflict = current["stockManaged"] and abs(qty_diff) > 0.00001
        has_conflict = has_conflict or abs(cost_diff) > 0.00001 or abs(price_diff) > 0.00001
        if has_conflict:
            conflicts.append(
                {
                    "当前商品": current["name"],
                    "真实商品": source["name"],
                    "当前数量": qty(current["quantity"]),
                    "真实数量": qty(source["quantity"]),
                    "当前成本": money(current["costPrice"]),
                    "真实成本均价": money(source["avgCost"]),
                    "当前售价": money(current["salePrice"]),
                    "真实预设售价": money(source["presetPrice"] or source["retailPrice"]),
                }
            )

    missing_rows = [
        {
            "商品编号": row["code"],
            "商品名称": row["name"],
            "类型": row["type"],
            "数量": qty(row["quantity"]),
            "成本均价": money(row["avgCost"]),
            "预设售价": money(row["presetPrice"] or row["retailPrice"]),
        }
        for row in source_missing_in_current
    ]

    extra_rows = [
        {
            "当前SKU": current["skuId"],
            "当前商品": current["name"],
            "类型": current["type"],
            "当前数量": qty(current["quantity"]),
            "最接近真实商品": source["name"] if source else "",
            "相似度": f"{score:.2f}",
        }
        for current, source, score in weak_matches
    ]
    likely_matches = [
        {
            "当前商品": current["name"],
            "真实商品": source["name"] if source else "",
            "相似度": f"{score:.2f}",
            "当前数量": qty(current["quantity"]),
            "真实数量": qty(source["quantity"]) if source else "",
            "当前成本": money(current["costPrice"]),
            "真实成本均价": money(source["avgCost"]) if source else "",
            "当前售价": money(current["salePrice"]),
            "真实预设售价": money(source["presetPrice"] or source["retailPrice"]) if source else "",
        }
        for current, source, score in weak_matches
        if source and score >= 0.6
    ]

    summary = {
        "source_count": len(source_rows),
        "source_stock_count": len(stock_source_rows),
        "source_total_quantity": sum(row["quantity"] for row in stock_source_rows),
        "source_total_value": sum(row["stockValue"] for row in stock_source_rows),
        "current_count": len(current_rows),
        "current_stock_quantity": sum(row["quantity"] for row in current_rows if row["stockManaged"]),
        "current_stock_value": sum(row["quantity"] * row["costPrice"] for row in current_rows if row["stockManaged"]),
        "direct_match_count": len(matches),
        "conflict_count": len(conflicts),
        "missing_count": len(missing_rows),
        "extra_count": len(extra_rows),
        "source_by_type": dict(source_by_type),
        "current_by_type": dict(current_by_type),
    }

    report = []
    report.append("# 库存状况列表对齐报告")
    report.append("")
    report.append("## 汇总")
    report.append("")
    report.append(f"- 真实库存表商品数：{summary['source_count']}")
    report.append(f"- 真实库存表中需要管理库存的商品数：{summary['source_stock_count']}")
    report.append(f"- 真实库存总数量：{qty(summary['source_total_quantity'])}")
    report.append(f"- 真实库存总成本金额：{money(summary['source_total_value'])}")
    report.append(f"- 当前系统商品数：{summary['current_count']}")
    report.append(f"- 当前系统实物库存总数量：{qty(summary['current_stock_quantity'])}")
    report.append(f"- 当前系统实物库存成本金额：{money(summary['current_stock_value'])}")
    report.append(f"- 名称直接/强匹配商品数：{summary['direct_match_count']}")
    report.append(f"- 已匹配但数量或价格冲突数：{summary['conflict_count']}")
    report.append(f"- 真实库存有、当前系统缺失数：{summary['missing_count']}")
    report.append(f"- 当前系统有、真实库存未强匹配数：{summary['extra_count']}")
    report.append("")
    report.append("## 真实库存类型分布")
    report.append("")
    report.append(table([{"类型": k, "商品数": v, "库存数量": qty(source_qty_by_type[k])} for k, v in sorted(source_by_type.items())], ["类型", "商品数", "库存数量"]))
    report.append("")
    report.append("## 当前系统类型分布")
    report.append("")
    report.append(table([{"类型": k, "商品数": v, "库存数量": qty(current_qty_by_type[k])} for k, v in sorted(current_by_type.items())], ["类型", "商品数", "库存数量"]))
    report.append("")
    report.append("## 已匹配但有冲突")
    report.append("")
    report.append(table(conflicts, ["当前商品", "真实商品", "当前数量", "真实数量", "当前成本", "真实成本均价", "当前售价", "真实预设售价"]) if conflicts else "无")
    report.append("")
    report.append("## 疑似可人工合并的商品")
    report.append("")
    report.append(table(likely_matches, ["当前商品", "真实商品", "相似度", "当前数量", "真实数量", "当前成本", "真实成本均价", "当前售价", "真实预设售价"]) if likely_matches else "无")
    report.append("")
    report.append("## 当前系统多出的商品或模拟商品")
    report.append("")
    report.append(table(extra_rows, ["当前SKU", "当前商品", "类型", "当前数量", "最接近真实商品", "相似度"]) if extra_rows else "无")
    report.append("")
    report.append("## 真实库存有、当前系统缺失的商品")
    report.append("")
    report.append(table(missing_rows, ["商品编号", "商品名称", "类型", "数量", "成本均价", "预设售价"]) if missing_rows else "无")
    report.append("")

    REPORT_MD.write_text("\n".join(report), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(str(REPORT_MD))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
